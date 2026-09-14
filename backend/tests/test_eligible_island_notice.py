"""Official eligible island notice: attachment discovery, xlsx parsing, change detection."""

from __future__ import annotations

import hashlib
from datetime import UTC, datetime
from io import BytesIO

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker

import app.models  # noqa: F401
from app.db.base import Base
from app.models import EligibleIslandCatalogSnapshot, EligibleIslandSnapshotEntry
from app.repositories.eligible_islands import CATALOG_KEY_ISLAND_VISIT_2026
from app.services.eligible_island_notice import (
    EligibleIslandNoticeError,
    ParsedIsland,
    SourceDocument,
    collect_eligible_island_catalog,
    fetch_eligible_island_notice_snapshot,
    parse_eligible_island_xlsx,
    source_fingerprint,
)

NOTICE_URL = "https://www.visitisland.kr/notice/12"
FETCHED_AT = datetime(2026, 9, 14, tzinfo=UTC)


def make_xlsx_bytes(*, headers: list[str], rows: list[list[object]], leading_rows: int = 0) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for _ in range(leading_rows):
        sheet.append(["2026 섬 방문의 해 대상 섬"])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class FakeResponse:
    def __init__(self, *, content: bytes = b"", text: str = "", status: int = 200, url: str = "") -> None:
        self.content = content or text.encode("utf-8")
        self.text = text
        self.status_code = status
        self.url = url  # final URL after redirects; filled in by make_http_get when empty

    def raise_for_status(self) -> None:
        if self.status_code >= 400:
            raise RuntimeError(f"http {self.status_code}")


def make_http_get(pages: dict[str, FakeResponse], *, calls: list[str] | None = None):
    def http_get(url: str) -> FakeResponse:
        if calls is not None:
            calls.append(url)
        if url not in pages:
            return FakeResponse(status=404, url=url)
        response = pages[url]
        if not response.url:
            response.url = url
        return response

    return http_get


@pytest.fixture
def db() -> Session:
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(bind=engine)
    TestingSessionLocal = sessionmaker(bind=engine, expire_on_commit=False)
    with TestingSessionLocal() as session:
        yield session


# --- parsing -----------------------------------------------------------------


def test_parse_xlsx_requires_island_and_jurisdiction_columns() -> None:
    payload = make_xlsx_bytes(headers=["섬명", "시군구"], rows=[["가거도", "전남 신안군"]])
    assert parse_eligible_island_xlsx(payload, filename="eligible.xlsx").entries == [
        ParsedIsland("가거도", "가거도", "전남 신안군")
    ]


def test_parse_xlsx_rejects_missing_jurisdiction_column() -> None:
    payload = make_xlsx_bytes(headers=["섬명"], rows=[["가거도"]])
    with pytest.raises(EligibleIslandNoticeError, match="parser_changed"):
        parse_eligible_island_xlsx(payload, filename="eligible.xlsx")


def test_parse_xlsx_joins_split_region_columns_and_finds_header_below_title_rows() -> None:
    payload = make_xlsx_bytes(
        headers=["연번", "시도", "시군구", "도서명"],
        rows=[[1, "전라남도", "신안군", " 가거도 "], [2, "경상남도", "통영시", "가거도"]],
        leading_rows=2,
    )
    parsed = parse_eligible_island_xlsx(payload, filename="south.xlsx")
    assert parsed.entries == [
        ParsedIsland("가거도", "가거도", "전라남도 신안군"),
        ParsedIsland("가거도", "가거도", "경상남도 통영시"),
    ]


def test_parse_xlsx_normalizes_nfc_and_whitespace_only() -> None:
    decomposed = "가거도".encode("utf-8").decode("utf-8")
    import unicodedata

    nfd_name = unicodedata.normalize("NFD", decomposed)
    payload = make_xlsx_bytes(headers=["섬명", "시군구"], rows=[[f"  {nfd_name}\t", "전남  신안군"]])
    parsed = parse_eligible_island_xlsx(payload, filename="eligible.xlsx")
    assert parsed.entries == [ParsedIsland(nfd_name.strip(), "가거도", "전남 신안군")]


def test_parse_xlsx_dedupes_same_name_and_jurisdiction_but_keeps_other_jurisdiction() -> None:
    payload = make_xlsx_bytes(
        headers=["섬명", "시군구"],
        rows=[["가거도", "전남 신안군"], ["가거도", "전남 신안군"], ["가거도", "경남 통영시"]],
    )
    parsed = parse_eligible_island_xlsx(payload, filename="eligible.xlsx")
    assert [(item.normalized_name, item.jurisdiction_name) for item in parsed.entries] == [
        ("가거도", "전남 신안군"),
        ("가거도", "경남 통영시"),
    ]


def test_parse_xlsx_rejects_non_zip_html_and_zero_entries() -> None:
    with pytest.raises(EligibleIslandNoticeError, match="parser_changed"):
        parse_eligible_island_xlsx(b"<html><body>login</body></html>", filename="eligible.xlsx")
    with pytest.raises(EligibleIslandNoticeError, match="parser_changed"):
        parse_eligible_island_xlsx(b"PK\x03\x04garbage", filename="eligible.xlsx")
    with pytest.raises(EligibleIslandNoticeError, match="parser_changed"):
        parse_eligible_island_xlsx(make_xlsx_bytes(headers=["섬명", "시군구"], rows=[]), filename="e.xlsx")
    with pytest.raises(EligibleIslandNoticeError, match="parser_changed"):
        parse_eligible_island_xlsx(make_xlsx_bytes(headers=["섬명", "시군구"], rows=[["", "전남 신안군"]]), filename="e.xlsx")


def test_source_fingerprint_is_order_independent() -> None:
    a = SourceDocument("https://x/a.xlsx", "a.xlsx", "1" * 64)
    b = SourceDocument("https://x/b.xlsx", "b.xlsx", "2" * 64)
    assert source_fingerprint([a, b]) == source_fingerprint([b, a])
    assert source_fingerprint([a]) != source_fingerprint([a, b])


# --- notice fetch --------------------------------------------------------------


def test_fetch_notice_snapshot_downloads_same_host_xlsx_attachments_only() -> None:
    south = make_xlsx_bytes(headers=["섬명", "시군구"], rows=[["가거도", "전남 신안군"]])
    east = make_xlsx_bytes(headers=["섬명", "시군구"], rows=[["울릉도", "경북 울릉군"]])
    html = (
        "<html><head><title>2026 대상 섬 목록 안내</title></head><body>"
        '<a href="/files/south.xlsx">전남</a>'
        '<a href="https://www.visitisland.kr/files/east.xlsx">경북</a>'
        '<a href="https://evil.example/files/west.xlsx">외부</a>'
        '<a href="/files/notes.pdf">안내문</a>'
        "</body></html>"
    )
    http_get = make_http_get(
        {
            NOTICE_URL: FakeResponse(text=html),
            "https://www.visitisland.kr/files/south.xlsx": FakeResponse(content=south),
            "https://www.visitisland.kr/files/east.xlsx": FakeResponse(content=east),
        }
    )
    snapshot = fetch_eligible_island_notice_snapshot(NOTICE_URL, http_get=http_get)
    assert snapshot.notice_title == "2026 대상 섬 목록 안내"
    assert [doc.filename for doc in snapshot.documents] == ["south.xlsx", "east.xlsx"]
    assert snapshot.documents[0].sha256 == hashlib.sha256(south).hexdigest()
    assert {item.normalized_name for item in snapshot.entries} == {"가거도", "울릉도"}
    assert snapshot.fingerprint == source_fingerprint(snapshot.documents)


SHEET_ID = "1Wx48HNm_acr3konWB5wTVSGupKBdnjKn4Qm4lXSnha8"
SHEET_EXPORT_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"


def test_fetch_notice_snapshot_follows_shortener_to_google_sheet_export() -> None:
    # Real site shape (2026-09-14): no attachment, only a buly.kr short link to a public Google Sheet.
    sheet = make_xlsx_bytes(headers=["연번", "시도", "시군구", "섬명"], rows=[[1, "인천광역시", "강화군", "주문도"]], leading_rows=3)
    html = '<html><head><title>프로모션 2차</title></head><body>※ 하기 리스트 <a href="https://buly.kr/8piNoSv">https://buly.kr/8piNoSv</a></body></html>'
    calls: list[str] = []
    http_get = make_http_get(
        {
            NOTICE_URL: FakeResponse(text=html),
            "https://buly.kr/8piNoSv": FakeResponse(text="<html>sheet</html>", url=f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit?usp=sharing"),
            SHEET_EXPORT_URL: FakeResponse(content=sheet),
        },
        calls=calls,
    )
    snapshot = fetch_eligible_island_notice_snapshot(NOTICE_URL, http_get=http_get)
    assert [(doc.url, doc.filename) for doc in snapshot.documents] == [(SHEET_EXPORT_URL, f"google-sheet-{SHEET_ID}.xlsx")]
    assert snapshot.documents[0].sha256 == hashlib.sha256(sheet).hexdigest()
    assert [(item.normalized_name, item.jurisdiction_name) for item in snapshot.entries] == [("주문도", "인천광역시 강화군")]
    assert calls == [NOTICE_URL, "https://buly.kr/8piNoSv", SHEET_EXPORT_URL]


def test_fetch_notice_snapshot_accepts_direct_google_sheet_link_once() -> None:
    sheet = make_xlsx_bytes(headers=["섬명", "시군구"], rows=[["가거도", "전남 신안군"]])
    link = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit#gid=0"
    html = f'<html><body><a href="{link}">리스트</a><a href="{link}">again</a></body></html>'
    http_get = make_http_get({NOTICE_URL: FakeResponse(text=html), SHEET_EXPORT_URL: FakeResponse(content=sheet)})
    snapshot = fetch_eligible_island_notice_snapshot(NOTICE_URL, http_get=http_get)
    assert [doc.url for doc in snapshot.documents] == [SHEET_EXPORT_URL]


def test_fetch_notice_snapshot_prefers_same_host_xlsx_over_list_links() -> None:
    south = make_xlsx_bytes(headers=["섬명", "시군구"], rows=[["가거도", "전남 신안군"]])
    html = f'<html><body><a href="/files/south.xlsx">전남</a><a href="https://buly.kr/8piNoSv">리스트</a></body></html>'
    calls: list[str] = []
    http_get = make_http_get(
        {NOTICE_URL: FakeResponse(text=html), "https://www.visitisland.kr/files/south.xlsx": FakeResponse(content=south)},
        calls=calls,
    )
    snapshot = fetch_eligible_island_notice_snapshot(NOTICE_URL, http_get=http_get)
    assert [doc.filename for doc in snapshot.documents] == ["south.xlsx"]
    assert "https://buly.kr/8piNoSv" not in calls


@pytest.mark.parametrize(
    ("href", "resolved"),
    [
        ("https://evil.example/files/list.xlsx", None),  # foreign host xlsx: never downloaded
        ("https://buly.kr/other", "https://evil.example/list.xlsx"),  # shortener to a non-sheet target
        ("https://bit.ly/8piNoSv", f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit"),  # shortener not on the allowlist
    ],
    ids=["foreign_xlsx", "shortener_to_non_sheet", "unlisted_shortener"],
)
def test_fetch_notice_snapshot_ignores_links_outside_the_allowlist(href: str, resolved: str | None) -> None:
    pages = {NOTICE_URL: FakeResponse(text=f'<html><body><a href="{href}">list</a></body></html>')}
    if resolved is not None:
        pages[href] = FakeResponse(text="<html></html>", url=resolved)
    calls: list[str] = []
    with pytest.raises(EligibleIslandNoticeError, match="parser_changed"):
        fetch_eligible_island_notice_snapshot(NOTICE_URL, http_get=make_http_get(pages, calls=calls))
    assert all(not url.endswith(".xlsx") or url.startswith("https://www.visitisland.kr/") for url in calls)
    assert "https://bit.ly/8piNoSv" not in calls


def test_fetch_notice_snapshot_reports_download_failed() -> None:
    html = '<html><body><a href="/files/south.xlsx">전남</a></body></html>'
    http_get = make_http_get({NOTICE_URL: FakeResponse(text=html)})
    with pytest.raises(EligibleIslandNoticeError, match="download_failed"):
        fetch_eligible_island_notice_snapshot(NOTICE_URL, http_get=http_get)
    with pytest.raises(EligibleIslandNoticeError, match="download_failed"):
        fetch_eligible_island_notice_snapshot(NOTICE_URL, http_get=make_http_get({}))


def test_fetch_notice_snapshot_without_attachments_is_parser_changed() -> None:
    http_get = make_http_get({NOTICE_URL: FakeResponse(text="<html><body>no files</body></html>")})
    with pytest.raises(EligibleIslandNoticeError, match="parser_changed"):
        fetch_eligible_island_notice_snapshot(NOTICE_URL, http_get=http_get)


# --- collection into pending snapshots ---------------------------------------------


def _pages(rows: list[list[object]]) -> dict[str, FakeResponse]:
    return {
        NOTICE_URL: FakeResponse(text='<html><head><title>대상 섬</title></head><body><a href="/files/list.xlsx">목록</a></body></html>'),
        "https://www.visitisland.kr/files/list.xlsx": FakeResponse(content=make_xlsx_bytes(headers=["섬명", "시군구"], rows=rows)),
    }


def test_collect_creates_pending_snapshot_with_entries(db: Session) -> None:
    result = collect_eligible_island_catalog(
        db, catalog_key=CATALOG_KEY_ISLAND_VISIT_2026, notice_url=NOTICE_URL, fetched_at=FETCHED_AT,
        http_get=make_http_get(_pages([["가거도", "전남 신안군"], ["홍도", "전남 신안군"]])),
    )
    assert result.outcome == "created"
    snapshot = db.get(EligibleIslandCatalogSnapshot, result.snapshot_id)
    assert snapshot is not None
    assert (snapshot.review_status, snapshot.entry_count, snapshot.notice_url) == ("pending", 2, NOTICE_URL)
    assert snapshot.attachment_filename == "list.xlsx"
    assert snapshot.attachment_documents == [
        {"url": "https://www.visitisland.kr/files/list.xlsx", "filename": "list.xlsx", "sha256": hashlib.sha256(
            make_xlsx_bytes(headers=["섬명", "시군구"], rows=[["가거도", "전남 신안군"], ["홍도", "전남 신안군"]])
        ).hexdigest()}
    ]
    names = db.scalars(select(EligibleIslandSnapshotEntry.normalized_name).where(EligibleIslandSnapshotEntry.snapshot_id == snapshot.id)).all()
    assert sorted(names) == ["가거도", "홍도"]


def test_collect_same_fingerprint_skips_parse_and_writes_nothing(db: Session, monkeypatch) -> None:
    from app.services import eligible_island_notice

    pages = _pages([["가거도", "전남 신안군"]])
    first = collect_eligible_island_catalog(
        db, catalog_key=CATALOG_KEY_ISLAND_VISIT_2026, notice_url=NOTICE_URL, fetched_at=FETCHED_AT, http_get=make_http_get(pages)
    )
    assert first.outcome == "created"

    def must_not_parse(*_args, **_kwargs):
        raise AssertionError("same fingerprint must not be parsed again")

    monkeypatch.setattr(eligible_island_notice, "parse_eligible_island_xlsx", must_not_parse)
    second = collect_eligible_island_catalog(
        db, catalog_key=CATALOG_KEY_ISLAND_VISIT_2026, notice_url=NOTICE_URL, fetched_at=FETCHED_AT, http_get=make_http_get(pages)
    )
    assert (second.outcome, second.snapshot_id) == ("unchanged", first.snapshot_id)
    assert db.scalar(select(EligibleIslandCatalogSnapshot.id).where(EligibleIslandCatalogSnapshot.id != first.snapshot_id)) is None


def test_collect_new_fingerprint_supersedes_older_pending(db: Session) -> None:
    first = collect_eligible_island_catalog(
        db, catalog_key=CATALOG_KEY_ISLAND_VISIT_2026, notice_url=NOTICE_URL, fetched_at=FETCHED_AT,
        http_get=make_http_get(_pages([["가거도", "전남 신안군"]])),
    )
    second = collect_eligible_island_catalog(
        db, catalog_key=CATALOG_KEY_ISLAND_VISIT_2026, notice_url=NOTICE_URL, fetched_at=FETCHED_AT,
        http_get=make_http_get(_pages([["가거도", "전남 신안군"], ["홍도", "전남 신안군"]])),
    )
    assert second.outcome == "created"
    assert db.get(EligibleIslandCatalogSnapshot, first.snapshot_id).review_status == "superseded"
    assert db.get(EligibleIslandCatalogSnapshot, second.snapshot_id).review_status == "pending"


def test_collect_failures_write_no_snapshot(db: Session) -> None:
    failed = collect_eligible_island_catalog(
        db, catalog_key=CATALOG_KEY_ISLAND_VISIT_2026, notice_url=NOTICE_URL, fetched_at=FETCHED_AT, http_get=make_http_get({})
    )
    assert (failed.outcome, failed.snapshot_id) == ("download_failed", None)
    changed = collect_eligible_island_catalog(
        db, catalog_key=CATALOG_KEY_ISLAND_VISIT_2026, notice_url=NOTICE_URL, fetched_at=FETCHED_AT,
        http_get=make_http_get(_pages([])),
    )
    assert (changed.outcome, changed.snapshot_id) == ("parser_changed", None)
    assert changed.error
    assert db.scalar(select(EligibleIslandCatalogSnapshot.id)) is None
