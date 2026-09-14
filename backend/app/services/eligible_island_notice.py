"""Official eligible island notice: find xlsx attachments, fingerprint, parse, stage pending snapshots.

The attachment bytes are never stored — only URL, filename, SHA-256 and parsed rows.
"""

from __future__ import annotations

import hashlib
import re
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from datetime import datetime
from html.parser import HTMLParser
from io import BytesIO
from typing import Any, Protocol
from urllib.parse import urljoin, urlparse

import httpx
from sqlalchemy.orm import Session

from app.repositories import eligible_islands as repository
from app.repositories.eligible_islands import normalize_island_name
from app.services.travelmonth_live_collector import DEFAULT_HEADERS, DEFAULT_TIMEOUT_SECONDS

PARSER_VERSION = "xlsx-v1"

# Only explicitly known header labels; anything else is a parser_changed failure.
ISLAND_NAME_HEADERS = {"섬명", "섬 명", "섬이름", "섬 이름", "도서명", "도서 명", "대상섬", "대상 섬", "섬"}
JURISDICTION_HEADERS = {"시군구", "시·군·구", "시/군/구", "시군", "관할", "관할지역", "관할 지역", "지자체", "소재지", "기초"}
REGION_HEADERS = {"시도", "시·도", "시/도", "광역", "광역시도", "광역 시도"}
_HEADER_SCAN_ROWS = 10
_ZIP_MAGIC = b"PK\x03\x04"


class EligibleIslandNoticeError(Exception):
    """Message is the outcome code: download_failed | parser_changed."""


@dataclass(frozen=True)
class ParsedIsland:
    display_name: str
    normalized_name: str
    jurisdiction_name: str
    raw_region_text: str | None = field(default=None, compare=False)


@dataclass(frozen=True)
class SourceDocument:
    url: str
    filename: str
    sha256: str

    def to_json(self) -> dict[str, str]:
        return {"url": self.url, "filename": self.filename, "sha256": self.sha256}


@dataclass(frozen=True)
class ParsedWorkbook:
    filename: str
    entries: list[ParsedIsland]


@dataclass(frozen=True)
class NoticeSnapshot:
    notice_url: str
    notice_title: str | None
    documents: list[SourceDocument]
    entries: list[ParsedIsland]
    fingerprint: str


@dataclass(frozen=True)
class EligibleIslandCollectionResult:
    outcome: str  # created | unchanged | download_failed | parser_changed
    snapshot_id: int | None = None
    entry_count: int = 0
    error: str | None = None


class _Response(Protocol):
    content: bytes
    text: str

    def raise_for_status(self) -> Any: ...


HttpGet = Callable[[str], _Response]


def _default_http_get(url: str) -> httpx.Response:
    return httpx.get(url, timeout=DEFAULT_TIMEOUT_SECONDS, follow_redirects=True, headers=DEFAULT_HEADERS)


def source_fingerprint(documents: Iterable[SourceDocument]) -> str:
    body = "\n".join(sorted(f"{item.url}|{item.sha256}" for item in documents))
    return hashlib.sha256(body.encode("utf-8")).hexdigest()


# --- xlsx parsing -------------------------------------------------------------------


def parse_eligible_island_xlsx(payload: bytes, *, filename: str) -> ParsedWorkbook:
    if not payload.startswith(_ZIP_MAGIC):
        raise EligibleIslandNoticeError("parser_changed") from ValueError(f"{filename}: not an xlsx workbook")
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:  # corrupt / encrypted / not a workbook
        raise EligibleIslandNoticeError("parser_changed") from exc

    entries: list[ParsedIsland] = []
    seen: set[tuple[str, str]] = set()
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        columns = _find_header(rows, filename=filename)
        if columns is None:
            continue
        name_col, jurisdiction_col, region_col = columns
        for row in rows:
            cells = [_cell_text(value) for value in row]
            if not any(cells):
                continue
            name = cells[name_col] if name_col < len(cells) else ""
            jurisdiction = cells[jurisdiction_col] if jurisdiction_col < len(cells) else ""
            region = cells[region_col] if region_col is not None and region_col < len(cells) else ""
            if not name or not jurisdiction:
                raise EligibleIslandNoticeError("parser_changed") from ValueError(
                    f"{filename}: blank island name or jurisdiction in a data row"
                )
            jurisdiction_name = normalize_island_name(f"{region} {jurisdiction}")
            key = (normalize_island_name(name), jurisdiction_name)
            if key in seen:
                continue
            seen.add(key)
            entries.append(
                ParsedIsland(
                    display_name=name,
                    normalized_name=key[0],
                    jurisdiction_name=jurisdiction_name,
                    raw_region_text=" ".join(part for part in (region, jurisdiction) if part) or None,
                )
            )
    if not entries:
        raise EligibleIslandNoticeError("parser_changed") from ValueError(f"{filename}: no island rows")
    return ParsedWorkbook(filename=filename, entries=entries)


def _find_header(rows: Iterable[tuple[Any, ...]], *, filename: str) -> tuple[int, int, int | None] | None:
    for index, row in enumerate(rows):
        if index >= _HEADER_SCAN_ROWS:
            break
        labels = [normalize_island_name(_cell_text(value)).replace(" ", "") for value in row]
        name_col = _first_index(labels, {label.replace(" ", "") for label in ISLAND_NAME_HEADERS})
        if name_col is None:
            continue
        jurisdiction_col = _first_index(labels, {label.replace(" ", "") for label in JURISDICTION_HEADERS})
        if jurisdiction_col is None:
            raise EligibleIslandNoticeError("parser_changed") from ValueError(
                f"{filename}: island column found but no jurisdiction column"
            )
        region_col = _first_index(labels, {label.replace(" ", "") for label in REGION_HEADERS})
        return name_col, jurisdiction_col, region_col
    return None


def _first_index(labels: list[str], wanted: set[str]) -> int | None:
    for index, label in enumerate(labels):
        if label in wanted:
            return index
    return None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


# --- notice page → attachments ---------------------------------------------------------


class _NoticeHtml(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.hrefs: list[str] = []
        self.title_parts: list[str] = []
        self._in_title = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "a":
            href = dict(attrs).get("href")
            if href:
                self.hrefs.append(href)
        elif tag == "title":
            self._in_title = True

    def handle_endtag(self, tag: str) -> None:
        if tag == "title":
            self._in_title = False

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self.title_parts.append(data)


def fetch_eligible_island_notice_snapshot(notice_url: str, *, http_get: HttpGet = _default_http_get) -> NoticeSnapshot:
    title, documents, payloads = _fetch_documents(notice_url, http_get=http_get)
    return NoticeSnapshot(
        notice_url=notice_url,
        notice_title=title,
        documents=documents,
        entries=_parse_documents(documents, payloads),
        fingerprint=source_fingerprint(documents),
    )


def _parse_documents(documents: list[SourceDocument], payloads: list[bytes]) -> list[ParsedIsland]:
    entries: list[ParsedIsland] = []
    seen: set[tuple[str, str]] = set()
    for document, payload in zip(documents, payloads, strict=True):
        for item in parse_eligible_island_xlsx(payload, filename=document.filename).entries:
            key = (item.normalized_name, item.jurisdiction_name)
            if key not in seen:
                seen.add(key)
                entries.append(item)
    return entries


def _fetch_documents(notice_url: str, *, http_get: HttpGet) -> tuple[str | None, list[SourceDocument], list[bytes]]:
    # Parsing is deferred so an unchanged fingerprint never opens a workbook.
    try:
        response = http_get(notice_url)
        response.raise_for_status()
        html = response.text
    except Exception as exc:
        raise EligibleIslandNoticeError("download_failed") from exc
    page = _NoticeHtml()
    page.feed(html)
    title = normalize_island_name("".join(page.title_parts)) or None
    host = urlparse(notice_url).netloc
    urls: list[str] = []
    for href in page.hrefs:
        absolute = urljoin(notice_url, href)
        parsed = urlparse(absolute)
        if parsed.netloc == host and parsed.path.lower().endswith(".xlsx") and absolute not in urls:
            urls.append(absolute)
    if not urls:
        raise EligibleIslandNoticeError("parser_changed") from ValueError("notice has no same-host .xlsx attachment")
    documents: list[SourceDocument] = []
    payloads: list[bytes] = []
    for url in urls:
        try:
            response = http_get(url)
            response.raise_for_status()
            payload = response.content
        except Exception as exc:
            raise EligibleIslandNoticeError("download_failed") from exc
        documents.append(SourceDocument(url=url, filename=urlparse(url).path.rsplit("/", 1)[-1], sha256=hashlib.sha256(payload).hexdigest()))
        payloads.append(payload)
    return title, documents, payloads


# --- collection ----------------------------------------------------------------------


def collect_eligible_island_catalog(
    db: Session,
    *,
    catalog_key: str,
    notice_url: str,
    fetched_at: datetime,
    http_get: HttpGet = _default_http_get,
) -> EligibleIslandCollectionResult:
    catalog = repository.lock_catalog_row(db, catalog_key=catalog_key)
    try:
        title, documents, payloads = _fetch_documents(notice_url, http_get=http_get)
        fingerprint = source_fingerprint(documents)
        existing = repository.find_snapshot_by_fingerprint(db, catalog_id=catalog.id, fingerprint=fingerprint)
        if existing is not None:
            db.commit()
            return EligibleIslandCollectionResult(outcome="unchanged", snapshot_id=existing.id, entry_count=existing.entry_count)
        entries = _parse_documents(documents, payloads)
    except EligibleIslandNoticeError as exc:
        db.rollback()
        return EligibleIslandCollectionResult(outcome=str(exc), error=_cause_text(exc))

    from app.services.eligible_island_catalog import stage_snapshot  # local: catalog service imports this module

    staged = stage_snapshot(
        db,
        catalog_key=catalog_key,
        entries=entries,
        notice_url=notice_url,
        notice_title=title,
        documents=documents,
        fetched_at=fetched_at,
    )
    db.commit()
    if staged.snapshot is None:
        # identical | suspicious_shrink (empty cannot happen: parser rejects zero rows)
        return EligibleIslandCollectionResult(outcome=staged.outcome, entry_count=len(entries))
    return EligibleIslandCollectionResult(outcome="created", snapshot_id=staged.snapshot.id, entry_count=staged.snapshot.entry_count)


def _cause_text(exc: BaseException) -> str:
    cause = exc.__cause__
    return f"{exc}: {cause}" if cause else str(exc)
